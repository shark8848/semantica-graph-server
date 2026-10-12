"""从 docs/功能清单.md 解析三列表格，导出为带样式的 Excel（一次性脚本）。"""
from __future__ import annotations

import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

MD = Path("docs/功能清单.md")
OUT = Path("docs/功能清单.xlsx")

# ---- 解析 markdown 表格（跳过表头分隔行与 > 引用 / # 标题 / --- 分隔）----
# 单元格内的竖线必须写成 `\|`（跳过转义竖线分割）；若一行仍被拆出非 3 列，说明表格写法有问题，
# 直接报错而不是静默丢行。
rows: list[list[str]] = []
for lineno, line in enumerate(MD.read_text(encoding="utf-8").splitlines(), start=1):
    s = line.strip()
    if not s.startswith("|"):
        continue
    cells = [c.strip() for c in re.split(r"(?<!\\)\|", s.strip("|"))]
    if len(cells) != 3:
        raise SystemExit(f"{MD}:{lineno} 表格行列数={len(cells)}（应为 3），单元格内竖线请写成 \\|")
    if all(re.fullmatch(r":?-{2,}:?", c) for c in cells):  # 分隔行
        continue
    rows.append(cells)

header, data = rows[0], rows[1:]


def clean(text: str) -> str:
    """去掉 markdown 残留标记：``code``、**粗体**与转义竖线。"""
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
    text = re.sub(r"`(.+?)`", r"\1", text)
    return text.replace("\\|", "|")


data = [[clean(cell) for cell in row] for row in data]

# ---- 构建工作簿 ----
wb = Workbook()
ws = wb.active
ws.title = "功能清单"

thin = Side(style="thin", color="B0B0B0")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
header_fill = PatternFill("solid", fgColor="1F4E79")
header_font = Font(bold=True, color="FFFFFF", size=11)
zebra_fill = PatternFill("solid", fgColor="F2F7FC")

ws.append(header)
for c, cell in enumerate(ws[1], start=1):
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = border

# 按一级功能分组做斑马色（组间交替）
prev_group = None
band = False
for r, cells in enumerate(data, start=2):
    group = cells[0]
    if group != prev_group:
        band = not band
        prev_group = group
    ws.append(cells)
    for c in range(1, 4):
        cell = ws.cell(row=r, column=c)
        cell.border = border
        cell.alignment = Alignment(vertical="top", wrap_text=(c == 3))
        if band:
            cell.fill = zebra_fill
    ws.cell(row=r, column=1).font = Font(bold=True)

widths = (30, 26, 90)
for i, w in enumerate(widths, start=1):
    ws.column_dimensions[get_column_letter(i)].width = w

ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:C{ws.max_row}"

wb.save(OUT)
print(f"导出完成：{OUT}  数据行={len(data)}  一级功能分组={len(set(r[0] for r in data))}")
